import sys
sys.stdout.reconfigure(encoding='utf-8')
import os
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.oxml import parse_xml

def set_bullet(paragraph, bullet_char="•", font_name="Arial"):
    pPr = paragraph._p.get_or_add_pPr()
    for child in list(pPr):
        if child.tag.endswith(('buNone', 'buAutoNum', 'buChar', 'buFont', 'buSzPct', 'buSzPts', 'buClr')):
            pPr.remove(child)
    buFont = parse_xml(f'<a:buFont xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="{font_name}"/>')
    buChar = parse_xml(f'<a:buChar xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" char="{bullet_char}"/>')
    pPr.append(buFont)
    pPr.append(buChar)

def set_no_bullet(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    for child in list(pPr):
        if child.tag.endswith(('buNone', 'buAutoNum', 'buChar', 'buFont', 'buSzPct', 'buSzPts', 'buClr')):
            pPr.remove(child)
    buNone = parse_xml('<a:buNone xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>')
    pPr.append(buNone)

def patch_deck(pptx_path):
    prs = Presentation(pptx_path)
    total_slides = len(prs.slides)
    print(f"Patching deck: {pptx_path} ({total_slides} slides)")

    # 1. Slide 8: Title font size 28pt so it fits on 1 line
    if total_slides >= 8:
        s8 = prs.slides[7]
        for sh in s8.shapes:
            if sh.has_text_frame and (sh.top / 914400.0 * 72) < 50:
                for p in sh.text_frame.paragraphs:
                    p.font.size = Pt(28)

    # 2. Slide 2: Table of Contents - Clean numbered list
    if total_slides >= 2:
        s2 = prs.slides[1]
        for sh in s2.shapes:
            if sh.has_text_frame and 50 < (sh.top / 914400.0 * 72) < 200:
                tf = sh.text_frame
                for p in tf.paragraphs:
                    set_no_bullet(p)
                    p.font.name = "Arial"
                    p.font.size = Pt(18)
                    p.space_after = Pt(14)

    # 3. Slide 3: Đặt vấn đề - Exactly 4 bullet points, NO empty paragraph, NO lone '-'
    if total_slides >= 3:
        s3 = prs.slides[2]
        for sh in s3.shapes:
            top_pt = sh.top / 914400.0 * 72
            if sh.has_text_frame and 50 < top_pt < 250:
                tf = sh.text_frame
                # Clear all paragraphs and rebuild
                tf.clear()
                
                items = [
                    [("77 triệu xe máy", True), (", chiếm 85–90% lưu lượng ⇒ ở Việt Nam ", False), ("biển hai dòng là đa số", True), (", không phải ngoại lệ", False)],
                    [("Biển xe mô tô chỉ ", False), ("190 × 140 mm", True), (", tỉ lệ khung hình ", False), ("1,357", True), (" — vừa là đối tượng nhỏ, vừa là bố cục hai dòng", False)],
                    [("Mọi bộ nhận dạng ký tự dựng sẵn đều ", False), ("giả định văn bản nằm trên một dòng ngang", True), (" — giả định đó ", False), ("sai", True), (" với đa số biển số Việt Nam", False)],
                    [("Ảnh thực tế còn thêm: bề mặt phản quang gây ", False), ("chói cục bộ", True), (", biển bám bụi, cong vênh, chụp nghiêng, ngược sáng", False)]
                ]
                
                for idx, runs_data in enumerate(items):
                    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                    set_bullet(p, bullet_char="•", font_name="Arial")
                    p.space_after = Pt(12)
                    for text, is_bold in runs_data:
                        r = p.add_run()
                        r.text = text
                        r.font.name = "Arial"
                        r.font.size = Pt(18)
                        r.font.bold = is_bold

    # 4. Slide 9, 11, 17: Thoroughly remove extraneous shapes at bottom margin
    for s_idx in [9, 11, 17]:
        if total_slides >= s_idx:
            slide = prs.slides[s_idx - 1]
            shapes_to_remove = []
            for sh in slide.shapes:
                width_pt = sh.width / 914400.0 * 72
                height_pt = sh.height / 914400.0 * 72
                left_pt = sh.left / 914400.0 * 72
                top_pt = sh.top / 914400.0 * 72
                is_page_num = (width_pt < 50 and height_pt < 50 and left_pt < 50)
                if is_page_num:
                    continue
                
                # Slide 11: MUST NOT have any picture (only table)
                if s_idx == 11 and (sh.shape_type == 13 or "Picture" in sh.name):
                    shapes_to_remove.append(sh)
                    continue
                
                # Slide 9 and 17: remove any second picture or any shape protruding at bottom
                if top_pt > 450 or (top_pt + height_pt) > 520:
                    shapes_to_remove.append(sh)
            
            for sh in shapes_to_remove:
                sp = sh._element
                sp.getparent().remove(sp)
                print(f"  [Slide {s_idx}] Removed extraneous shape: {sh.name}")

    # 5. Slide 11: Table 'Bộ luật hậu xử lý' - O vs 0 in row 5
    if total_slides >= 11:
        s11 = prs.slides[10]
        for sh in s11.shapes:
            if sh.has_table:
                tbl = sh.table
                for row in tbl.rows:
                    for cell in row.cells:
                        for p in cell.text_frame.paragraphs:
                            p.font.name = "Arial"
                if len(tbl.rows) >= 5:
                    c2 = tbl.cell(4, 1)
                    c2.text = "O → 0 hợp lý, 0 → O không bao giờ — chiều đúng là 0 → D"
                    for p in c2.text_frame.paragraphs:
                        p.font.name = "Arial"
                        p.font.size = Pt(13)

    # 6. Slide 19: Ba kết quả khác với dự đoán ban đầu - clean bullets
    if total_slides >= 19:
        s19 = prs.slides[18]
        for sh in s19.shapes:
            top_pt = sh.top / 914400.0 * 72
            if sh.has_text_frame and 50 < top_pt < 250:
                for p in sh.text_frame.paragraphs:
                    if p.text.strip():
                        set_bullet(p, bullet_char="•", font_name="Arial")
                        p.font.name = "Arial"
                        p.font.size = Pt(17)
                        p.space_after = Pt(12)

    # 7. Slide 21: Table 'Hướng phát triển' - 6 distinct rows, synchronized with Table 5.3 in thesis
    if total_slides >= 21:
        s21 = prs.slides[20]
        for sh in s21.shapes:
            if sh.has_table:
                tbl = sh.table
                if len(tbl.rows) >= 7:
                    # Column 3 values from user & Table 5.3:
                    # Row 1: 1
                    # Row 2: 1
                    # Row 3: 3, 4
                    # Row 4: 2, 5
                    # Row 5: —
                    # Row 6: —
                    col3_vals = ["1", "1", "3, 4", "2, 5", "—", "—"]
                    for r_idx in range(1, 7):
                        c1 = tbl.cell(r_idx, 0)
                        c2 = tbl.cell(r_idx, 1)
                        c3 = tbl.cell(r_idx, 2)
                        
                        c3.text = col3_vals[r_idx - 1]
                        
                        for c in [c1, c2, c3]:
                            for p in c.text_frame.paragraphs:
                                p.font.name = "Arial"
                                p.font.size = Pt(12)
                        
                        for p in c1.text_frame.paragraphs:
                            p.alignment = PP_ALIGN.CENTER
                        for p in c2.text_frame.paragraphs:
                            p.alignment = PP_ALIGN.LEFT
                        for p in c3.text_frame.paragraphs:
                            p.alignment = PP_ALIGN.CENTER

    # 8. Slide 22: Kết luận chung - Exactly 4 bullet points, NO stray '0', NO leading blank line
    if total_slides >= 22:
        s22 = prs.slides[21]
        for sh in s22.shapes:
            top_pt = sh.top / 914400.0 * 72
            if sh.has_text_frame and 50 < top_pt < 250:
                tf = sh.text_frame
                tf.clear()
                
                items22 = [
                    [("Chạy đầu cuối trên CPU: phát hiện ", False), ("mAP@0,5 = 0,9829", True)],
                    [("Đột phá trên biển hai dòng nhờ biến đổi ảnh (thay vì đổi mô hình): tăng ", False), ("+34,92 điểm", True)],
                    [("Hậu xử lý theo vị trí ", False), ("+13,28 điểm", True), (", ", False), ("0 ca làm sai lệch", True), (" / 2.801 biển", False)],
                    [("Điểm nghẽn còn lại: ", False), ("biển hai dòng", True), (" (S₁ = 0,7234)", False)]
                ]
                
                for idx, runs_data in enumerate(items22):
                    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                    set_bullet(p, bullet_char="•", font_name="Arial")
                    p.space_after = Pt(14)
                    for text, is_bold in runs_data:
                        r = p.add_run()
                        r.text = text
                        r.font.name = "Arial"
                        r.font.size = Pt(18)
                        r.font.bold = is_bold

    # 9. Slide 23: Cảm ơn & Q&A
    if total_slides >= 23:
        s23 = prs.slides[22]
        for sh in s23.shapes:
            top_pt = sh.top / 914400.0 * 72
            if sh.has_text_frame and 70 < top_pt < 200:
                tf = sh.text_frame
                for p in tf.paragraphs:
                    set_no_bullet(p)
                    p.alignment = PP_ALIGN.CENTER
                if len(tf.paragraphs) >= 5:
                    p1 = tf.paragraphs[0]
                    p1.font.size = Pt(19)
                    p1.font.bold = True
                    p1.font.color.rgb = RGBColor(0, 51, 128)
                    p1.space_before = Pt(10)
                    p1.space_after = Pt(18)

                    p2 = tf.paragraphs[1]
                    p2.font.size = Pt(15)
                    p2.font.bold = False
                    p2.space_after = Pt(8)

                    p3 = tf.paragraphs[2]
                    p3.text = "Sinh viên thực hiện: Phạm Nguyễn Thế Châu | Nguyễn Công Hậu | Nguyễn Minh Hiếu | Phạm Công Thành"
                    p3.font.size = Pt(14)
                    p3.font.name = "Arial"
                    p3.font.bold = False
                    p3.space_after = Pt(24)

                    p4 = tf.paragraphs[3]
                    p4.font.size = Pt(17)
                    p4.font.bold = True
                    p4.font.italic = True
                    p4.font.color.rgb = RGBColor(34, 34, 34)
                    p4.space_after = Pt(16)

                    p5 = tf.paragraphs[4]
                    p5.font.size = Pt(24)
                    p5.font.bold = True
                    p5.font.color.rgb = RGBColor(0, 102, 204)
                    p5.space_before = Pt(10)

    prs.save(pptx_path)
    print(f"Successfully patched {pptx_path}")

    # Export to PDF via PowerPoint COM if available
    try:
        import win32com.client
        full_pptx = os.path.abspath(pptx_path)
        full_pdf = os.path.splitext(full_pptx)[0] + ".pdf"
        ppt = win32com.client.Dispatch("PowerPoint.Application")
        pres = ppt.Presentations.Open(full_pptx, False, False, False)
        pres.SaveAs(full_pdf, 32)
        pres.Close()
        ppt.Quit()
        print(f"Successfully exported PDF: {full_pdf}")
    except Exception as e:
        print(f"[warn] Could not export PDF via PowerPoint COM: {e}")

if __name__ == "__main__":
    deck = sys.argv[1] if len(sys.argv) > 1 else "docs/slides/12-slides-mon-hoc.pptx"
    patch_deck(deck)

